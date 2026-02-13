import os
from typing import Any
from langchain_text_splitters import RecursiveCharacterTextSplitter
from app.core.config import settings


class DocumentProcessor:
    """Extracts text from documents and splits into chunks for RAG."""

    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
            length_function=len,
            separators=["\n\n", "\n", ". ", " ", ""],
        )

    async def process(
        self, document_id: str, file_path: str, file_name: str
    ) -> dict[str, Any]:
        """Process a document and return extracted text chunks."""
        ext = os.path.splitext(file_name)[1].lower()

        if ext == ".pdf":
            text, pages = self._extract_pdf(file_path)
        elif ext == ".txt":
            text, pages = self._extract_txt(file_path)
        elif ext == ".docx":
            text, pages = self._extract_docx(file_path)
        elif ext == ".pptx" or ext == ".ppt":
            text, pages = self._extract_pptx(file_path)
        elif ext == ".xlsx" or ext == ".xls":
            text, pages = self._extract_xlsx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

        # Split into chunks
        chunks = self.text_splitter.split_text(text)

        return {
            "document_id": document_id,
            "text": text,
            "pages": pages,
            "chunks": chunks,
        }

    def _extract_pdf(self, file_path: str) -> tuple[str, int]:
        from pypdf import PdfReader

        reader = PdfReader(file_path)
        text = ""
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n\n"
        return text.strip(), len(reader.pages)

    def _extract_txt(self, file_path: str) -> tuple[str, int]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        pages = max(1, len(text) // 3000)
        return text.strip(), pages

    def _extract_docx(self, file_path: str) -> tuple[str, int]:
        from docx import Document

        doc = Document(file_path)
        text = "\n\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        pages = max(1, len(text) // 3000)
        return text.strip(), pages

    def _extract_pptx(self, file_path: str) -> tuple[str, int]:
        from pptx import Presentation

        prs = Presentation(file_path)
        slides_text = []
        for i, slide in enumerate(prs.slides, 1):
            slide_content = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_content.append(text)
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            slide_content.append(row_text)
            if slide_content:
                slides_text.append(f"[Slide {i}]\n" + "\n".join(slide_content))
        text = "\n\n".join(slides_text)
        return text.strip(), len(prs.slides)

    def _extract_xlsx(self, file_path: str) -> tuple[str, int]:
        from openpyxl import load_workbook

        wb = load_workbook(file_path, read_only=True, data_only=True)
        sheets_text = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows_text = []
            for row in ws.iter_rows(values_only=True):
                cell_values = [str(c) if c is not None else "" for c in row]
                row_str = " | ".join(v for v in cell_values if v)
                if row_str:
                    rows_text.append(row_str)
            if rows_text:
                sheets_text.append(f"[Sheet: {sheet_name}]\n" + "\n".join(rows_text))
        wb.close()
        text = "\n\n".join(sheets_text)
        pages = max(1, len(wb.sheetnames))
        return text.strip(), pages
